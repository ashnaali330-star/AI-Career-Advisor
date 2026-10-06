from io import BytesIO
from pathlib import Path
import re


def parse_resume(filename: str, content: bytes) -> str:
    """Extract text from PDF, DOCX, or TXT upload bytes."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        from docx import Document
        doc = Document(BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    elif ext == ".txt":
        text = content.decode("utf-8-sig", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 30:
        raise ValueError("Very little text was extracted. The file may be empty or a scanned PDF without selectable text.")
    return text
